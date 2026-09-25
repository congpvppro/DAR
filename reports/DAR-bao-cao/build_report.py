"""Build the Vietnamese DAR report from frozen local evidence (no inference)."""
from pathlib import Path
import os, json, hashlib, collections, zipfile, io
os.environ.setdefault('MPLCONFIGDIR', '/tmp/dar-report-mpl')
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
ASSETS = OUT / 'assets'
SOURCES = {}
def read_json(rel):
    p = ROOT / rel
    SOURCES[rel] = hashlib.sha256(p.read_bytes()).hexdigest()
    return json.loads(p.read_text())
def vn(x, digits=2):
    return f'{x:,.{digits}f}'.replace(',', '_').replace('.', ',').replace('_', '.')
def integer(x): return vn(x, 0)

splits = {}
for sp in ['train', 'test']:
    rel = f'data/DAR-R1-annotations/{sp}.jsonl'
    p = ROOT / rel
    SOURCES[rel] = hashlib.sha256(p.read_bytes()).hexdigest()
    rows = [json.loads(l) for l in p.read_text().splitlines() if l.strip()]
    seq = [json.loads(r['conversations'][1]['value'])['segments'] for r in rows]
    for r, ss in zip(rows, seq):
        assert ss and abs(ss[0]['start_time']) < 1e-6
        assert abs(ss[-1]['end_time'] - r['video_duration']) < 1e-5
        assert all(s['end_time'] > s['start_time'] for s in ss)
        assert all(abs(a['end_time'] - b['start_time']) < 1e-5 for a, b in zip(ss, ss[1:]))
    splits[sp] = {'rows': rows, 'seq': seq}

allseq = splits['train']['seq'] + splits['test']['seq']
trans = collections.Counter((a['emotion'], b['emotion']) for ss in allseq for a,b in zip(ss, ss[1:]))
outgoing = collections.Counter()
for (a,b), n in trans.items(): outgoing[a] += n
dwell = collections.defaultdict(list)
for ss in allseq:
    for s in ss: dwell[s['emotion']].append(s['end_time']-s['start_time'])
durations = np.array([d for ds in dwell.values() for d in ds])
stats = {'splits': {}, 'transition_count': sum(trans.values()),
         'same_emotion_adjacent': sum(n for (a,b),n in trans.items() if a==b),
         'duration_mean': float(durations.mean()),
         'duration_quantiles': dict(zip(['p25','p50','p75','p90','p95'], map(float,np.quantile(durations,[.25,.5,.75,.9,.95])))),
         'duration_le5': float(np.mean(durations<=5)), 'duration_le10': float(np.mean(durations<=10)),
         'per_emotion_duration': [{'emotion':k,'n':len(v),'mean':float(np.mean(v)),'median':float(np.median(v))} for k,v in sorted(dwell.items())],
         'transitions': [{'from':a,'to':b,'count':n,'outgoing':outgoing[a],'conditional':n/outgoing[a]} for (a,b),n in trans.most_common()]}
for sp, v in splits.items():
    ds = [s['end_time']-s['start_time'] for ss in v['seq'] for s in ss]
    stats['splits'][sp] = {'videos':len(v['rows']),'segments':len(ds), 'transitions':sum(len(ss)-1 for ss in v['seq']),
       'mean_video_duration':float(np.mean([r['video_duration'] for r in v['rows']])),
       'mean_segment_duration':float(np.mean(ds)), 'median_segment_duration':float(np.median(ds)),
       'segment_count_distribution':dict(collections.Counter(map(len,v['seq'])))}
assert len(durations)==36908 and sum(trans.values())==21821
assert stats['same_emotion_adjacent']==0
full = read_json('research/gap-study-20260907/full-error-breakdown.json')
baseline = read_json('research/gap-study-20260907/analysis.json')
dar = read_json('outputs/paper-repro-gpu01/final-summary.json')
qwen = read_json('outputs/baselines-gpu0123/qwen25-vl-3b/summary.json')
study = read_json('research/hypothesis-validation-20260908/analysis.json')
audit = read_json('research/original-grounding-audit-20260909/summary.json')
panel = full['panels']['0.5']
assert full['segmentation_counts'] == {'under':583,'over':288,'equal':570}
assert panel['qualified_pairs']==1550 and panel['correct']==483
assert audit['n']==64 and audit['event_error_videos']==27

plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10, 'text.color':'black',
  'axes.labelcolor':'black','axes.edgecolor':'black','xtick.color':'black','ytick.color':'black',
  'axes.spines.top':False,'axes.spines.right':False, 'savefig.facecolor':'white'})
def save_fig(fig, name):
    path = ASSETS / f'{name}.png'
    fig.savefig(path, dpi=220, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    # Exact monochrome pixels, including all lettering, outlines and patterns.
    with Image.open(path) as im:
        im.convert('L').point(lambda p: 255 if p>=190 else 0, mode='1').save(path)
    return path

top = trans.most_common(12)
fig, ax = plt.subplots(figsize=(10,4.8))
ax.barh([f'{a} → {b}' for (a,b),n in top], [n for _,n in top], facecolor='white',edgecolor='black',hatch='///',height=.62)
for i,(_,n) in enumerate(top): ax.text(n+5,i,str(n),va='center')
ax.invert_yaxis();ax.set_xlim(0,490);ax.set_xlabel('Số lần chuyển giữa hai đoạn kề nhau (train + test)')
fig.tight_layout(); save_fig(fig,'01-transitions')

fig, ax = plt.subplots(figsize=(9,4.5))
for sp,style,label in [('train','-','Train (33.195 đoạn)'),('test','--','Test (3.713 đoạn)')]:
    d=np.sort([s['end_time']-s['start_time'] for ss in splits[sp]['seq'] for s in ss])
    ax.step(d,100*np.arange(1,len(d)+1)/len(d),where='post',color='black',linestyle=style,linewidth=1.6,label=label)
ax.set_xlim(0,48);ax.set_ylim(0,102);ax.set_xlabel('Thời lượng duy trì một nhãn cảm xúc (giây)');ax.set_ylabel('Tỷ lệ đoạn có thời lượng ≤ t (%)')
ax.legend(loc='lower right',frameon=False);fig.tight_layout();save_fig(fig,'02-persistence')

mistakes=[r for r in panel['confusion'] if r['gt']!=r['pred']][:10]
fig,ax=plt.subplots(figsize=(10,4.6))
ax.barh([f'{r["gt"]} → {r["pred"]}' for r in mistakes],[r['count'] for r in mistakes],facecolor='white',edgecolor='black',hatch='///',height=.6)
for i,r in enumerate(mistakes):ax.text(r['count']+.3,i,f'{r["count"]}/{r["qualified_gt"]} ({100*r["count"]/r["qualified_gt"]:.1f}%)',va='center',fontsize=9)
ax.invert_yaxis();ax.set_xlim(0,27);ax.set_xlabel('Số cặp đoạn sai nhãn; mẫu số là số cặp đủ IoU của nhãn GT')
fig.tight_layout();save_fig(fig,'03-confusions')

fig,axs=plt.subplots(1,2,figsize=(10,4),gridspec_kw={'width_ratios':[1,1.6]})
counts=[583,570,288]
axs[0].bar(['Chia thiếu','Đúng số','Chia thừa'],counts,color='white',edgecolor='black',hatch='///')
for i,n in enumerate(counts):axs[0].text(i,n+12,f'{n}\n({100*n/1441:.1f}%)',ha='center',fontsize=9)
axs[0].set_ylim(0,730);axs[0].set_ylabel('Số video');axs[0].tick_params(axis='x',labelsize=9)
mat=np.zeros((6,6),dtype=int)
for r in full['segment_count_matrix']: mat[r['gt_count']-1,r['predicted_count']-1]=r['videos']
axs[1].set_xlim(.5,6.5);axs[1].set_ylim(6.5,.5)
axs[1].set_xticks(range(1,7));axs[1].set_yticks(range(1,7))
for i in range(6):
    for j in range(6):
        axs[1].text(j+1,i+1,str(mat[i,j]) if mat[i,j] else '·',ha='center',va='center',fontweight='bold' if i==j else 'normal')
for x in np.arange(.5,7):
    axs[1].axvline(x,color='black',linewidth=.4);axs[1].axhline(x,color='black',linewidth=.4)
axs[1].set_xlabel('Số đoạn dự đoán');axs[1].set_ylabel('Số đoạn GT')
fig.tight_layout();save_fig(fig,'04-segmentation')

fig,ax=plt.subplots(figsize=(9,3.5))
labels=['DAR • phát hành','DAR • thận trọng','Qwen • phát hành','Qwen • thận trọng']
rates=[47,48,0,0]
ax.barh(labels,[100*x/64 for x in rates],color='white',edgecolor='black',hatch='///',height=.55)
for i,(n,amb) in enumerate(zip(rates,[4,5,21,19])):ax.text(100*n/64+1,i,f'{n}/64; mơ hồ: {amb}/64',va='center',fontsize=9)
ax.invert_yaxis();ax.set_xlim(0,110);ax.set_xticks([0,25,50,75,100]);ax.set_xlabel('Output có ít nhất một lỗi sự kiện rõ trên ảnh lặp (%)')
fig.tight_layout();save_fig(fig,'05-static-grounding')

# Document: simple A4 pages, black typography, white tables, black rules.
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8)
sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.header_distance=Cm(.7);sec.footer_distance=Cm(.7)
for st in doc.styles:
    if st.type in (1,2,3):
        st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        if st._element.rPr is not None:
            for c in st._element.rPr.findall(qn('w:color')):
                for a in ['themeColor','themeTint','themeShade']:c.attrib.pop(qn('w:'+a),None)
normal=doc.styles['Normal'];normal.font.size=Pt(11.5)
normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.12
for name,size in [('Title',17),('Heading 1',14),('Heading 2',12),('Heading 3',11.5)]:
    st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True
    st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(10)
doc.styles['Caption'].font.italic=True
doc.styles['Caption'].paragraph_format.space_after=Pt(7)
header=sec.header.paragraphs[0];header.text='DAR • BÁO CÁO PHÂN TÍCH VÀ THỰC NGHIỆM';header.runs[0].font.size=Pt(9)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=foot.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
def p(text,style=None):
    para=doc.add_paragraph(text,style);para.paragraph_format.widow_control=True;return para
def h(text,level=2):return doc.add_heading(text,level=level)
def page():doc.add_page_break()
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,txt in zip(t.rows[0].cells,headers):c.text=str(txt)
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
    pr=t._tbl.tblPr
    bd=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+name);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'000000');bd.append(el)
    pr.append(bd)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
        if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,cell in enumerate(row.cells):
            if widths:cell.width=Cm(widths[j])
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(3);para.paragraph_format.line_spacing=1.0
                for run in para.runs:run.font.size=Pt(10.5);run.font.bold=i==0
    p('').paragraph_format.space_after=Pt(0)
    return t
def figure(name,caption,width=16.5):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next=True
    para.add_run().add_picture(str(ASSETS/f'{name}.png'),width=Cm(width))
    p(caption,'Caption')

# Page 1
h('BÁO CÁO PHÂN TÍCH BÀI BÁO DAR',0)
p('Benchmarking Dynamic Affective Reasoning: A Viewer-Centric Video Emotion Dataset')
p('Bài báo: Zhiyan Zhang và cộng sự • arXiv:2607.10238v1, 11/07/2026\nDữ liệu thực nghiệm tổng hợp đến 09/09/2026. [1–7]')
p('DAR nghiên cứu cảm xúc mà video gây ra cho người xem theo thời gian. Với mỗi video, mô hình phải xác định ranh giới các pha cảm xúc, gán một trong 27 nhãn và giải thích sự kiện thị giác dẫn đến cảm xúc đó. DAR-R1 sử dụng Qwen2.5-VL-3B, được tác giả huấn luyện bằng SFT rồi GRPO. [1]')
p('Báo cáo này tổng hợp các thử nghiệm đã thực hiện trong dự án: tái lập benchmark, phân tích lỗi, đối chứng đầu vào và kiểm tra bằng chứng trong lời giải thích. Thống kê dữ liệu được tính lại từ annotations; không chạy thêm suy luận mô hình để lập báo cáo.')
h('1. THỐNG KÊ DỮ LIỆU',1)
h('1.1. Quy mô và cấu trúc tập dữ liệu')
rows=[]
for label,key in [('Số video','videos'),('Số đoạn cảm xúc','segments'),('Số chuyển tiếp giữa các đoạn','transitions')]:
    a=stats['splits']['train'][key];b=stats['splits']['test'][key];rows.append([label,integer(a),integer(b),integer(a+b)])
rows.extend([['Số đoạn/video',vn(33195/13646),vn(3713/1441),vn(36908/15087)],
 ['Thời lượng video TB (giây)',vn(stats['splits']['train']['mean_video_duration']),vn(stats['splits']['test']['mean_video_duration']),vn(np.mean([r['video_duration'] for v in splits.values() for r in v['rows']]))],
 ['Thời lượng đoạn TB (giây)','5,86','5,68','5,84'],['Số lớp cảm xúc','27','27','27']])
table(['Thống kê','Train','Test','Toàn bộ'],rows,[8,3,3,3])
p('Nguồn: train.jsonl và test.jsonl công khai trong dự án. Mỗi chuyển tiếp là một cặp đoạn liền kề trong cùng video; không nối đoạn giữa hai video. [2]','Caption')
p('Dữ liệu tập trung vào video ngắn có nhiều pha: số đoạn phổ biến nhất là 2 ở train và 3 ở test. Test có 109/1.441 video một đoạn (7,56%). Khác biệt này cần được ghi nhận khi dùng cấu trúc phân đoạn học từ train làm đối chứng.')

# Page 2
page();h('1.2. Phân bố chuyển cảm xúc (emotion transition distribution)')
p('Đếm N(a → b) trên tất cả các cặp đoạn liền kề. Báo cáo hai đại lượng: tần suất tuyệt đối N(a → b), và xác suất có điều kiện P(b | a) = N(a → b) / Σc N(a → c). Mẫu số của P(b | a) chỉ gồm đoạn mang nhãn a còn có đoạn kế tiếp; đoạn cuối video không đóng góp.')
figure('01-transitions','Hình 1. Mười hai chuyển tiếp phổ biến nhất, tính từ 21.821 chuyển tiếp của 15.087 video train + test. Đây là chuyển tiếp nhãn GT, không phải nhầm lẫn của mô hình. [2]')
table(['Chuyển tiếp','Số lần / số lần rời nhãn nguồn','P(đích | nguồn)'],[[f'{a} → {b}',f'{trans[a,b]} / {outgoing[a]}',vn(100*trans[a,b]/outgoing[a])+'%'] for a,b in [('Craving','Satisfaction'),('Anxiety','Fear'),('Surprise','Amusement')]], [6.7,6.8,3.5])
p('Craving → Satisfaction đạt 46,85% và Anxiety → Fear đạt 30,06%, khớp các giá trị làm tròn 0,47 và 0,30 trong Hình 3(c) của bài báo. Surprise → Amusement có số lần xuất hiện lớn nhất (437), nhưng “nhiều lần nhất” không đồng nghĩa “xác suất có điều kiện cao nhất”. [1, 2]')
p('Rút ra: nhãn dữ liệu có các hướng chuyển tiếp thường gặp, tạo cơ sở cho đối chứng dùng prior cảm xúc. Đây là mô tả annotations; riêng tần suất không chứng minh cơ chế nhân quả tâm lý của người xem.')

# Page 3
page();h('1.3. Thời gian duy trì cảm xúc (emotion persistence)')
p('Trong báo cáo, persistence là thời lượng một pha mang cùng nhãn: d = end_time − start_time. Đây là thời gian duy trì được ghi trong video, không phải phép đo cảm xúc sinh lý liên tục hoặc thời gian cảm xúc tồn tại sau khi clip kết thúc.')
p('Cả 21.821 cặp đoạn kề nhau đều khác nhãn. Bài báo mô tả việc gộp các đoạn liên tiếp cùng cảm xúc; vì vậy đường chéo ma trận chuyển tiếp bằng 0 là hệ quả của cách phân đoạn, không có nghĩa cảm xúc không được duy trì. [1, 2]')
figure('02-persistence','Hình 2. Phân phối tích lũy thời lượng đoạn: trục đứng là tỷ lệ đoạn có thời lượng không vượt quá t. Nét liền: train; nét đứt: test. [2]')
table(['Đại lượng, toàn bộ 36.908 đoạn','Giá trị'],[
 ['Trung bình / trung vị','5,84 / 4,50 giây'],['Khoảng tứ phân vị (P25–P75)','2,50–7,80 giây'],
 ['Phân vị 90% / 95%','12,00 / 15,00 giây'],['Đoạn dài ≤5 giây / ≤10 giây','56,22% / 84,83%'],
 ['Surprise: trung bình; số đoạn','3,70 giây; n = 2.653'],['Sadness: trung bình; số đoạn','8,52 giây; n = 705']], [11,6])
p('Rút ra: đa số pha cảm xúc chỉ kéo dài vài giây; trung bình lớn hơn trung vị cho thấy phân bố lệch về phía thời lượng dài. Surprise có thời lượng trung bình ngắn nhất, Sadness dài nhất trong 27 nhãn. Khác biệt này mô tả tập video đã chọn, chưa tách ảnh hưởng của nội dung, thời lượng clip và quy trình gán nhãn.')

# Page 4
page();h('2. THIẾT LẬP THỬ NGHIỆM (EXPERIMENT SETUP)',1)
h('2.1. Tái lập và chấm lại trên toàn bộ test')
p('Tôi sử dụng annotations chính thức, video VCE và checkpoint công khai; đánh giá đầu ra gồm start_time, end_time, emotion và reason. Prompt yêu cầu phân tích video im lặng theo cảm xúc người xem. Các thí nghiệm dưới đây dùng trọng số có sẵn; tôi không huấn luyện lại SFT hoặc GRPO. [3–5]')
table(['Thành phần','Cách thực hiện'],[
 ['Dữ liệu','1.441 video test; 3.713 đoạn GT. Đủ file video; hai shard 720 + 721 ghép lại khớp test JSONL từng byte.'],
 ['DAR-R1','Checkpoint đầy đủ BF16; chạy hai tiến trình trên GPU 0–1, mỗi GPU chứa một bản mô hình.'],
 ['Môi trường tái lập','Python 3.10.19; PyTorch 2.8.0+cu128; Transformers 4.57.1; vLLM 0.11.0; RTX 3090 24 GB.'],
 ['Tiền xử lý DAR-R1','Đặt use_fast=False (slow image processor). Run full-test dùng luồng video của script tái lập; giới hạn 16 frame là cấu hình của thí nghiệm can thiệp ở mục 2.2.'],
 ['Sinh đáp án','Seed 1234; temperature 0,1; top-p 0,9; batch 8; context 16.384; tối đa 4.096 token đầu ra.'],
 ['Xử lý lỗi DAR-R1','Giữ riêng run strict. Chỉ chạy lại 5 ID không parse được với batch 1, tối đa 8.192 token; sau đó đủ 1.441 output không rỗng.'],
 ['Đối chiếu baseline','Qwen2.5-VL-3B-Instruct dùng prompt DAR; sau retry còn 1 ID lỗi. AffectGPT là adapter frame-only local, không nhận audio hoặc transcript.'],
 ['Tái phân tích','Dùng output lịch sử, bao gồm retry; chấm lại cả ba mô hình và đối chứng thời lượng, không coi đây là inference mới.']], [4.2,12.8])
p('Các file strict được giữ nguyên; kết quả bổ sung sau retry được báo riêng vì ngân sách sinh đáp án đã đổi. Output parse được thành danh sách không rỗng chỉ là độ phủ output, không bảo đảm timeline hợp lệ hoặc phủ đủ video.')
h('2.2. Lịch sử các đợt thực nghiệm')
table(['Đợt','Phạm vi'],[
 ['03/09/2026','Tái lập DAR-R1 full-test và chạy các baseline.'],
 ['07/09/2026','Phân tích lại full-test; pilot mới 48 video × 4 điều kiện × 2 model = 384 lượt; thêm 4 lượt AffectGPT.'],
 ['08/09/2026','736 lượt chẩn đoán timing; chạy lại toàn bộ 736 lượt sau sửa FPS. Chỉ run timing-corrected được dùng cho kết quả kiểm chứng chính.'],
 ['09/09/2026','Audit lại reasoning của DAR-R1 trên 64 video gốc; không chạy inference mới.']], [3.2,13.8])

# Page 5
page();h('2.3. Thiết kế đối chứng để kiểm tra mô hình dùng bằng chứng gì')
p('Run kiểm chứng chính chọn 64 ID test mới so với pilot 48 ID trước đó, seed 20260908; clip dài (0; 30] giây, có ít nhất 16 frame mã hóa. Mẫu gồm 158 đoạn GT và không phân tầng theo số đoạn. Đây là mẫu dùng cho can thiệp, không phải tập test chưa từng được khảo sát. [6]')
table(['Đối chứng','Đầu vào và mục đích'],[
 ['Original / static','Cùng 16 frame lấy đều, hoặc lặp chính xác frame giữa 16 lần. Kiểm tra khả năng viện dẫn chuyển động khi đầu vào không thay đổi.'],
 ['Hai prompt','Prompt phát hành và prompt bổ sung yêu cầu không bịa sự kiện, chấp nhận một pha, phân biệt quan sát với suy đoán.'],
 ['Gray / gray → white','16 frame xám; hoặc 8 frame xám + 8 frame trắng. Lưới 12 thời lượng: 4, 6, 8, 10, 12, 15, 18, 20, 24, 30, 40, 60 giây.'],
 ['Factual probe','Hỏi riêng “các frame có thay đổi thị giác không?”, trả boolean và câu bằng chứng; không hỏi nhãn cảm xúc.'],
 ['Duration + train prior','Không xem frame; lấy mode train là 2 đoạn, chia đôi thời lượng và gán Interest → Amusement theo mode vị trí của train hai đoạn.']], [4.4,12.6])
p('Cả DAR-R1 và Qwen dùng greedy, seed 20260908, context 12.288, batch tối đa 8; 4.096 token cho cảm xúc, 256 cho factual. Requested pixel budget là 100.352; tensor tổng hợp thực tế 16 × 3 × 308 × 308. Dùng bốn RTX 3090, giữ nguyên trọng số và không retry chọn lọc.')
p('Ma trận chính: 512 lượt natural (64 × original/static × 2 prompt × 2 model), 48 lượt cảm xúc gray và 176 lượt factual; tổng 736. Adapter được sửa để truyền FPS lấy mẫu vào processor; run trước sửa chỉ dùng chẩn đoán, không gộp với run chính. [6]')
h('2.4. Chỉ số đánh giá và đơn vị đếm')
table(['Chỉ số','Định nghĩa dùng trong báo cáo'],[
 ['SC-Acc','Số video đúng số đoạn / số video đánh giá; ghi rõ mẫu số toàn test hay chỉ output hợp lệ.'],
 ['Index mIoU','Ghép đoạn dự đoán và GT cùng chỉ số tới min(Npred, NGT); lấy IoU trung bình theo evaluator phát hành.'],
 ['Emo-Acc có điều kiện','C / Q; Q là số cặp IoU ≥ 0,5, C là số cặp trong Q còn đúng nhãn.'],
 ['Joint recall theo index','C / G, với G là toàn bộ đoạn GT, kể cả GT không có cặp. Dùng để kiểm phần bị bỏ sót.'],
 ['Lỗi grounding','Tỷ lệ video/output có ≥1 khẳng định thị giác sai rõ. Không đồng nhất với sai nhãn cảm xúc hoặc chỉ sai timestamp.']], [4.5,12.5])

# Page 6
page();h('3. KẾT QUẢ QUAN SÁT VÀ NHẬN XÉT',1)
h('3.1. Kết quả tái lập trên toàn bộ benchmark')
table(['Nguồn / cấu hình','Output không rỗng','SC-Acc','mIoU','Emo-Acc*'],[
 ['DAR-R1 — bài báo','Không nêu','41,50%','52,30%','28,60%'],
 ['DAR-R1 — strict 4.096','1.436/1.441','39,69%','50,16%','31,09%'],
 ['DAR-R1 — sau retry','1.441/1.441','39,56%','50,10%','31,16%'],
 ['Qwen — bài báo','Không nêu','25,40%','41,70%','15,00%'],
 ['Qwen — sau retry','1.440/1.441','26,81%','40,39%','15,31%']], [6.2,3.2,2.5,2.5,2.6])
p('*Các dòng local dùng C/Q tại IoU ≥ 0,5; bài báo không mô tả đủ matching/mẫu số để bảo đảm hoàn toàn đồng nhất với cột Emo-Acc. SC-Acc của strict và Qwen trong bảng này dùng output hợp lệ làm mẫu số. [1, 3, 4]','Caption')
p('So với số công bố, DAR-R1 sau retry thấp hơn 1,94 điểm phần trăm về SC-Acc và 2,20 điểm về mIoU; accuracy nhãn có điều kiện local cao hơn 2,56 điểm so với con số Emo-Acc trong bài báo. Đây là đối chiếu mô tả, chưa phải tái lập chính xác toàn bộ giao thức tác giả.')
h('So sánh bằng mẫu số GT cố định')
table(['Model / đối chứng local','SC-Acc /1.441','Index mIoU','Joint recall /3.713 GT'],[
 ['DAR-R1','39,56%','50,10%','13,01%'],['Qwen2.5-VL-3B','26,79%','40,39%','5,36%'],
 ['AffectGPT frame-only','11,17%','28,07%','0,94%'],['Duration + train prior','37,40%','48,92%','4,47%']], [7,3.3,3,3.7])
p('AffectGPT có 1.418/1.441 output không rỗng; đây là nhánh adapter local, không phải tái lập tương đương hàng AffectGPT của bài báo. Qwen 26,79% ở bảng này dùng toàn bộ 1.441 video, khác 26,81% trên 1.440 output hợp lệ ở bảng trên. [5]','Caption')
p('DAR-R1 vẫn tốt hơn Qwen khi dùng mẫu số GT cố định: joint recall tăng 7,65 điểm phần trăm, bootstrap theo video 95% CI [6,28; 8,96]. Tuy nhiên, giá trị tuyệt đối 13,01% cho thấy còn nhiều đoạn chưa được nhận đúng cả thời gian lẫn cảm xúc.')
p('Không có GPT-Score thực chạy trong dự án. Bài báo có prompt judge, nhưng repository không cung cấp đầy đủ evaluator dịch vụ để tái lập nguyên trạng; vì vậy không gán điểm GPT Avg 3,3/5 của bài báo cho kết quả local.')

# Page 7
page();h('3.2. Lỗi nhãn cảm xúc: nhầm những gì?')
p('DAR-R1 tạo 3.349 đoạn cho 3.713 GT. Theo cách ghép cùng chỉ số, có 3.005 cặp được so; 1.550 cặp đạt IoU ≥ 0,5, trong đó 483 đúng nhãn và 1.067 sai nhãn (68,84%). Ngoài ra, 1.455 cặp không đạt IoU và 708 đoạn GT không có cặp cùng chỉ số. [5]')
figure('03-confusions','Hình 3. Mười hướng nhầm nhãn phổ biến nhất trên full-test; vẽ lại từ dữ liệu nguồn của error-visualizations. Chiều mũi tên là GT → dự đoán, không phải chuyển cảm xúc thực trong video. [5, 8]')
table(['Nhãn GT','Đúng / cặp đạt IoU','Tỷ lệ đúng có điều kiện'],[
 ['Awkwardness','7/72','9,72%'],['Fear','14/89','15,73%'],['Joy','11/48','22,92%'],
 ['Amusement','71/168','42,26%'],['Romance','31/43','72,09%']], [6.4,5.3,5.3])
p('Awkwardness thường bị đổi thành Amusement; Fear còn bị nhầm với Anxiety và Confusion. Nhầm nhãn vẫn xuất hiện khi thời gian khớp chặt: ở IoU ≥ 0,8 có 404/597 cặp sai (67,67%). Vì vậy, sai cảm xúc không chỉ có thể giải thích bằng lệch ranh giới.')
p('Rút ra: cần cải thiện phân biệt appraisal gần nhau và báo độ đúng theo từng nhãn. Các tỷ lệ trên chỉ xét cặp đã đủ overlap, không phải recall trên mọi GT. Craving chỉ có 5 cặp đủ IoU nên tỷ lệ 60% không đủ để kết luận đây là nhãn mô hình xử lý vững. Nhầm nhãn tự nó không chứng minh mô hình bịa sự kiện.')

# Page 8
page();h('3.3. Lỗi phân đoạn và giới hạn của mIoU')
figure('04-segmentation','Hình 4. Trái: chia thiếu, đúng số và chia thừa trên 1.441 video. Phải: số video theo số đoạn GT và số đoạn dự đoán; chữ đậm trên đường chéo là đúng số đoạn. [5, 8]')
p('DAR-R1 chia thiếu ở 583 video (40,46%), chia thừa ở 288 video (19,99%) và đúng số đoạn ở 570 video (39,56%). Với GT có 3 đoạn, 361/664 video bị dự đoán thành 2 đoạn. Ngược lại, ở 109 video GT một đoạn, mô hình chỉ đúng số ở 19 video và chia thừa 90 video.')
p('Quan sát này gợi ý mô hình thường đưa đầu ra về cấu trúc ít pha quen thuộc; riêng số đếm chưa xác định khuynh hướng đến từ dữ liệu SFT, GRPO, prompt hay decoding.')
table(['Đầu ra dùng để chấm','Index mIoU','Joint recall / mọi GT'],[
 ['DAR-R1 đầy đủ','50,10%','13,01%'],['Duration + train prior','48,92%','4,47%'],
 ['Chỉ giữ đoạn đầu của DAR-R1','59,07%','7,81%']], [9.2,3.6,4.2])
p('Đối chứng chỉ biết thời lượng đạt mIoU 48,92%, gần giá trị 50,10% của DAR-R1, nhưng nhận đúng cảm xúc và thời gian ít hơn rõ. Chỉ giữ đoạn đầu còn làm mIoU tăng lên 59,07% dù joint recall giảm còn 7,81%. Cách cắt này vi phạm yêu cầu phủ timeline; đây là phản ví dụ của cách đọc metric, không phải một cải tiến mô hình. [5]')
p('Rút ra: temporal overlap cao chưa đủ chứng minh hiểu diễn tiến cảm xúc. Khi báo kết quả cần đặt mIoU cạnh SC-Acc, joint recall/F1 và kiểm tra tính hợp lệ của timeline. Không suy ra DAR-R1 bỏ qua hình ảnh: mô hình vẫn hơn đối chứng ở khả năng đúng cả nhãn lẫn thời gian.')

# Page 9
page();h('3.4. Đối chứng ảnh lặp: nhiều pha có đi kèm sự kiện bị bịa?')
p('Run timing-corrected hoàn thành đủ 736 lượt. Tác vụ cảm xúc có 559/560 output không rỗng; factual có 176/176 JSON boolean hợp lệ. Có 3/559 output cảm xúc không rỗng vẫn chứa interval lỗi sau chuẩn hóa, nên parse thành công không đồng nghĩa timeline đúng. [6]')
figure('05-static-grounding','Hình 5. Lỗi sự kiện rõ trên 64 đầu vào lặp frame giữa cho từng model/prompt. Mẫu số giữ đủ 64 output; số mơ hồ được ghi riêng. Đây là kết quả dưới can thiệp static. [6, 8]')
p('DAR-R1 dùng prompt phát hành vẫn sinh ≥2 pha ở 58/64 đầu vào static (90,63%), nhưng tiêu chí lỗi chỉ tính khi lời giải thích khẳng định thay đổi thị giác không xảy ra. Theo tiêu chí này có 47/64 output lỗi rõ (73,44%); Wilson 95% CI [61,52%; 82,70%]. Nhiều pha tự nó không phải lỗi: cảm xúc người xem có thể thay đổi khi nhìn một ảnh lâu hơn.')
p('Ví dụ ID 51665: trên 16 bản sao giống hệt của một frame, DAR-R1 viết “its front end lifts sharply in a jarring, unnatural motion”, rồi kể xe dừng lại. Chuyển động đó không tồn tại trong tensor được đưa vào mô hình. [6]')
p('Prompt thận trọng không cho thấy giảm lỗi DAR-R1: 47/64 → 48/64; chênh +1,56 điểm phần trăm, CI95% [−7,81; 10,94], Holm p = 1,0. Qwen có 0 lỗi rõ trên static nhưng còn 21/64 và 19/64 ca mơ hồ; trên lưới gray, Qwen prompt phát hành lại có 12/12 output lỗi. Vì vậy không kết luận Qwen không hallucinate.')
p('Cách chấm: đọc 304 raw answers static/gray; đối chiếu bằng hai lượt AI che model/prompt trên 206 trường hợp, gồm mọi lỗi sơ bộ, mọi ca mơ hồ và một mẫu ca âm. Giữ bất đồng là mơ hồ. Đây không phải chấm bởi chuyên gia con người; CI không bao gồm sai số annotation.')

# Page 10
page();h('3.5. Audit lời giải thích trên video gốc')
p('Để kiểm tra lỗi có xuất hiện khi không can thiệp đầu vào hay không, tôi sử dụng lại 64 output DAR-R1 ở điều kiện original, prompt phát hành. Audit đọc toàn bộ reasoning, xem 16 frame mô hình nhận và frame nguồn lấy 2 fps + frame cuối; bổ sung frame dày ở đoạn cần xác minh. Tái dựng tensor đầu vào khớp hash ở 64/64 ca. [7]')
table(['Kết quả audit','Video /64','Tỷ lệ'],[
 ['Có ít nhất một lỗi hình ảnh rõ','38','59,38%'],['Trong 38 ca trên: có lỗi sự kiện','27','42,19%'],
 ['Chưa đủ bằng chứng kết luận','17','26,56%'],['Chưa thấy lỗi rõ','9','14,06%']], [10,3.5,3.5])
p('Ba nhóm chính 38 + 17 + 9 = 64; 27 ca lỗi sự kiện là tập con của 38 ca lỗi hình ảnh, không cộng thêm. Đơn vị là output có ít nhất một khẳng định sai, không phải tỷ lệ câu sai. [7]','Caption')
table(['ID và thời gian mô hình gán','Khẳng định trong output','Kết quả đối chiếu trong audit'],[
 ['01203; 1,5–4,0 giây','“tower suddenly collapses at 1.8 seconds”','Người đàn ông gạt đĩa khỏi chồng chai; tháp vẫn đứng, sau đó người này ăn mừng.'],
 ['05128; 0,0–11,5 giây','“he unexpectedly pulls out a pen and begins signing the paper”','Người đàn ông rút ghế gấp đen từ gói quà cạnh võ đài rồi giơ ghế.'],
 ['60095; 3,8–5,4 giây','“before landing smoothly on the tiled platform”','Người trượt ván mất thăng bằng và rơi lăn xuống bậc thang.']], [3.7,6.1,7.2])
p('Các câu trích lấy từ output local; mô tả đối chiếu kế thừa hồ sơ audit và ảnh có timestamp, không chép từ ground truth cảm xúc. Có thể tra toàn bộ 64 ca tại research/original-grounding-audit-20260909/review.html. [7]','Caption')
p('Rút ra: bằng chứng hiện có ủng hộ kết luận hẹp rằng DAR-R1 có thể dùng diễn biến hoặc kết quả thị giác sai làm tiền đề giải thích ngay trên video gốc. Lỗi không chỉ xuất hiện khi thay video bằng ảnh tĩnh hoặc màn xám.')
p('Giới hạn: mẫu 64 clip ngắn được chọn cho nghiên cứu can thiệp, không đại diện toàn bộ benchmark; một tác nhân AI thực hiện audit, chưa có người chấm độc lập. Không suy rộng 59,38% hoặc 42,19% thành tỷ lệ lỗi của 1.441 video. Chưa thấy lỗi rõ cũng không chứng minh toàn bộ reasoning đúng.')

# Page 11
page();h('3.6. Có thể kết luận đến đâu?')
p('Trên cùng 64 video gốc, DAR-R1 prompt phát hành đạt mIoU 57,38% và joint recall 15,19%; Qwen đạt 35,14% và 2,53%; đối chứng thời lượng đạt 52,44% và 6,33%. DAR-R1 có lợi ích thực về tác vụ. Kiểm định tương đương DAR với prior trong biên ±5 điểm mIoU chưa đạt: CI90% của chênh [−0,70; 10,87] không nằm trọn trong [−5; 5]. [6]')
p('Factual probe cho thấy cả hai model trả lời “không thay đổi” đúng boolean ở 64/64 ảnh lặp, nhưng đối chứng có thay đổi xám → trắng chỉ được DAR trả lời đúng 1/12 và Qwen 0/12. Probe có xu hướng trả lời false, nên không thể từ kết quả âm suy ra mô hình luôn nhìn đúng và chỉ sai bước suy luận cảm xúc. [6]')
table(['Kết luận được dữ liệu hỗ trợ','Điều chưa được chứng minh'],[
 ['DAR-R1 cải thiện so với Qwen ở benchmark và mẫu original.','Chưa xác định riêng phần cải thiện do SFT, GRPO hay cấu trúc annotations.'],
 ['Prior thời lượng tạo được temporal overlap đáng kể.','Chưa chứng minh prior tương đương hoặc thay thế được DAR-R1.'],
 ['Có lời giải thích viện dẫn sự kiện sai ở static và video gốc.','Chưa chứng minh cơ chế “chọn nhãn trước rồi bịa lý do”, hoặc tỷ lệ lỗi toàn benchmark.'],
 ['Prompt thận trọng đã thử không cho thấy giảm lỗi DAR.','Không suy ra mọi cách prompting đều không hiệu quả.']], [8.5,8.5])
h('3.7. Hướng nghiên cứu rút ra')
p('Ưu tiên đánh giá độ đúng của sự kiện được viện dẫn khi giải thích chuyển cảm xúc, đồng thời đo độ bao phủ các sự kiện/đoạn. Một nhãn hợp lý hoặc lời giải thích trôi chảy chỉ có giá trị khi tiền đề thị giác được xác minh.')
p('Thử nghiệm tiếp theo nên giữ sự kiện đích rồi thay lịch sử liên quan bằng lịch sử đối chứng; xóa đúng bằng chứng mô hình viện dẫn và so với xóa phần không liên quan. Cần người xem xác nhận can thiệp có thực sự đổi cảm xúc kỳ vọng, cho phép nhiều nhãn hợp lệ và bất định ranh giới. Đây là đề xuất, chưa được chạy.')
p('Khi đánh giá phương pháp mới, báo đồng thời số đoạn, temporal overlap, joint recall/F1, lỗi timeline và tỷ lệ output có sự kiện sai; giữ đối chứng thời lượng và đối chứng có/không có thay đổi thị giác. Chỉ quy kết tác động của SFT/GRPO sau ablation có dữ liệu, token và compute tương đương.')

# Page 12
page();h('NGUỒN VÀ KHẢ NĂNG ĐỐI CHIẾU',1)
p('Số trong bảng và hình dùng dấu phẩy thập phân. “Điểm phần trăm” là chênh lệch tuyệt đối giữa hai tỷ lệ phần trăm. Trừ các dòng ghi “bài báo”, kết quả được lấy từ hiện vật local hoặc tính lại từ annotations.')
refs=[
 ('[1] Bài báo gốc','Zhiyan Zhang và cộng sự. Benchmarking Dynamic Affective Reasoning: A Viewer-Centric Video Emotion Dataset. arXiv:2607.10238v1, 2026. Thống kê: §3.5, Hình 3; mô hình và thí nghiệm: §4–5, Table 2.\nhttps://arxiv.org/abs/2607.10238\nBản lưu: research/sources/DAR-2607.10238v1.pdf'),
 ('[2] Nhãn dữ liệu','data/DAR-R1-annotations/train.jsonl\ndata/DAR-R1-annotations/test.jsonl\nNguồn công khai: https://huggingface.co/datasets/aiaiaizzy/DAR-R1'),
 ('[3] Tái lập DAR-R1','outputs/paper-repro-gpu01/final-summary.json\noutputs/paper-repro-gpu01/paper-comparison.md\nresearch/run_logs/paper-repro-gpu01/preflight.txt'),
 ('[4] Baseline Qwen','outputs/baselines-gpu0123/qwen25-vl-3b/summary.json\noutputs/baselines-gpu0123/qwen25-vl-3b/paper-comparison.md'),
 ('[5] Tái phân tích full-test và lỗi','research/gap-study-20260907/analysis.json\nresearch/gap-study-20260907/full-error-breakdown.json\nresearch/gap-study-20260907/report.vi.md'),
 ('[6] Kiểm chứng với đầu vào có kiểm soát','research/hypothesis-validation-20260908/analysis.json\nresearch/hypothesis-validation-20260908/report.vi.md\nRaw outputs và protocol: thư mục timing-corrected/ cùng cấp.'),
 ('[7] Audit video gốc','research/original-grounding-audit-20260909/summary.json\nresearch/original-grounding-audit-20260909/report.vi.md\nresearch/original-grounding-audit-20260909/review.html'),
 ('[8] Nguồn trực quan hóa lỗi','research/error-visualizations/provenance.json\nHình 3–5 trong báo cáo được vẽ lại bằng đen trắng từ các file thống kê nguồn, giữ mẫu số và phạm vi phân tích.')]
for title,body in refs:
    para=p(title);para.runs[0].bold=True;para.paragraph_format.space_after=Pt(2)
    pp=p(body);pp.paragraph_format.space_after=Pt(5)
    for r in pp.runs:r.font.size=Pt(10)
p('Kèm theo báo cáo: dataset-statistics.json lưu số liệu tính lại; provenance.json lưu SHA-256 nguồn; build_report.py lưu mã tạo hình và Word. Mọi chữ, viền bảng và nét biểu đồ dùng màu đen trên nền trắng.','Caption')

for rel in ['research/sources/DAR-2607.10238v1.pdf','research/error-visualizations/provenance.json',
 'research/hypothesis-validation-20260908/report.vi.md','research/original-grounding-audit-20260909/report.vi.md',
 'research/run_logs/paper-repro-gpu01/preflight.txt','research/gap-study-20260907/report.vi.md']:
    SOURCES[rel]=hashlib.sha256((ROOT/rel).read_bytes()).hexdigest()

# Explicitly set all actual text black, including headers/footers and tables.
for root in [doc._element,sec.header._element,sec.footer._element]:
    for run in root.iter(qn('w:r')):
        rp=run.find(qn('w:rPr'))
        if rp is None:rp=OxmlElement('w:rPr');run.insert(0,rp)
        for old in rp.findall(qn('w:color')):rp.remove(old)
        col=OxmlElement('w:color');col.set(qn('w:val'),'000000');rp.append(col)
        lang=OxmlElement('w:lang');lang.set(qn('w:val'),'vi-VN');rp.append(lang)
doc.core_properties.title='Báo cáo phân tích và thực nghiệm DAR'
doc.core_properties.subject='Thống kê dữ liệu, thiết lập thử nghiệm và phân tích kết quả'
doc.core_properties.author=''
doc.core_properties.keywords='DAR, emotion transition, emotion persistence, experiment setup'
path=OUT/'Bao_cao_DAR.docx';doc.save(path)
# Remove latent colored styles/themes, so editing later stays monochrome too.
with zipfile.ZipFile(path) as z: parts={n:z.read(n) for n in z.namelist()}
for name,content in list(parts.items()):
    if name.endswith('.xml') and name.startswith('word/'):
        root=etree.fromstring(content)
        for el in root.iter(qn('w:color')):
            el.attrib.clear();el.set(qn('w:val'),'000000')
        for el in root.iter():
            if qn('w:color') in el.attrib:
                el.set(qn('w:color'),'000000')
                for attr in ['themeColor','themeTint','themeShade']:
                    el.attrib.pop(qn('w:'+attr),None)
        # The bundled default Title style contains a colored bottom rule.
        for el in list(root.iter(qn('w:pBdr'))):
            el.getparent().remove(el)
        if 'theme/' in name:
            ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
            for parent in root.findall('.//a:clrScheme/*',ns):
                for child in list(parent):parent.remove(child)
                child=etree.SubElement(parent,'{'+ns['a']+'}srgbClr')
                child.set('val','FFFFFF' if etree.QName(parent).localname.startswith('lt') else '000000')
        parts[name]=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
    for name,content in parts.items():z.writestr(name,content)
(OUT/'dataset-statistics.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2)+'\n')
(OUT/'provenance.json').write_text(json.dumps({'sources_sha256':SOURCES,'new_model_inference_runs':0,'report':path.name},ensure_ascii=False,indent=2)+'\n')
print(path)
print('Document paragraphs:',len(doc.paragraphs),'tables:',len(doc.tables),'images:',len(doc.inline_shapes))
